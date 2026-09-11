import os
from pathlib import Path
from typing import List, Dict, Any
from PIL import Image
import pytesseract

from langchain_core.documents import Document
from langchain_community.document_loaders import PyPDFLoader, TextLoader
from langchain_text_splitters import RecursiveCharacterTextSplitter

class UniversalDocumentLoader:
    """Handles parsing and ingestion of diverse document formats into LangChain Documents."""

    def __init__(self, chunk_size: int = 1000, chunk_overlap: int = 200):
        self.chunk_size = chunk_size
        self.chunk_overlap = chunk_overlap
        self.splitter = RecursiveCharacterTextSplitter(
            chunk_size=chunk_size,
            chunk_overlap=chunk_overlap,
            separators=["\n\n", "\n", " ", ""]
        )

    def load_file(self, file_path: str) -> List[Document]:
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")

        ext = path.suffix.lower().lstrip(".")
        file_name = path.name
        file_size_kb = round(path.stat().st_size / 1024, 1)

        raw_docs: List[Document] = []

        if ext == "pdf":
            loader = PyPDFLoader(str(path))
            raw_docs = loader.load()
            for d in raw_docs:
                d.metadata["source"] = file_name
                d.metadata["file_type"] = "PDF"
                d.metadata["file_size_kb"] = file_size_kb

        elif ext == "docx":
            try:
                import docx
                doc = docx.Document(str(path))
                full_text = []
                for p in doc.paragraphs:
                    if p.text.strip():
                        full_text.append(p.text)
                for table in doc.tables:
                    for row in table.rows:
                        row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                        if row_text:
                            full_text.append(row_text)

                content = "\n\n".join(full_text)
                if content.strip():
                    raw_docs.append(Document(
                        page_content=content,
                        metadata={"source": file_name, "file_type": "DOCX", "file_size_kb": file_size_kb}
                    ))
            except Exception as e:
                raise RuntimeError(f"Error parsing Word DOCX file: {e}")

        elif ext in ["csv", "xlsx", "xls"]:
            try:
                import pandas as pd
                if ext == "csv":
                    df = pd.read_csv(str(path))
                else:
                    df = pd.read_excel(str(path))

                # Format dataframe summary + row chunks
                summary = f"Table Data Summary for {file_name}:\nColumns: {list(df.columns)}\nTotal Rows: {len(df)}\n\nSample Data:\n{df.head(20).to_string()}"
                raw_docs.append(Document(
                    page_content=summary,
                    metadata={"source": file_name, "file_type": "Spreadsheet", "file_size_kb": file_size_kb}
                ))

                # Also include row batches for deeper search
                for i in range(0, len(df), 25):
                    batch_str = df.iloc[i:i+25].to_string()
                    raw_docs.append(Document(
                        page_content=f"--- {file_name} Rows {i+1}-{min(i+25, len(df))} ---\n{batch_str}",
                        metadata={"source": file_name, "file_type": "Spreadsheet", "file_size_kb": file_size_kb}
                    ))
            except Exception as e:
                raise RuntimeError(f"Error parsing tabular file: {e}")

        elif ext in ["txt", "md", "py", "js", "ts", "html", "css", "json", "java", "cpp", "c", "sql", "xml", "yaml", "yml"]:
            loader = TextLoader(str(path), encoding="utf-8", autodetect_encoding=True)
            loaded = loader.load()
            for d in loaded:
                d.metadata["source"] = file_name
                d.metadata["file_type"] = ext.upper()
                d.metadata["file_size_kb"] = file_size_kb
            raw_docs.extend(loaded)

        elif ext in ["png", "jpg", "jpeg", "webp", "bmp"]:
            try:
                img = Image.open(path)
                extracted_text = pytesseract.image_to_string(img)
                if extracted_text.strip():
                    raw_docs.append(Document(
                        page_content=extracted_text,
                        metadata={"source": file_name, "file_type": "Image OCR", "file_size_kb": file_size_kb}
                    ))
                else:
                    raise ValueError(f"No readable text found in image '{file_name}' via OCR.")
            except Exception as e:
                raise RuntimeError(f"OCR Error for image: {e}")

        else:
            # Fallback text read
            try:
                with open(path, "r", encoding="utf-8", errors="ignore") as f:
                    content = f.read()
                if content.strip():
                    raw_docs.append(Document(
                        page_content=content,
                        metadata={"source": file_name, "file_type": ext.upper(), "file_size_kb": file_size_kb}
                    ))
            except Exception as e:
                raise RuntimeError(f"Unsupported or unreadable file format '{ext}': {e}")

        return raw_docs

    def process_and_chunk(self, file_paths: List[str]) -> tuple[List[Document], List[Dict[str, Any]], List[str]]:
        """Processes multiple files, chunks them, and returns (chunks, loaded_meta, errors)."""
        all_chunks = []
        loaded_meta = []
        errors = []

        for p in file_paths:
            try:
                raw = self.load_file(p)
                if not raw:
                    errors.append(f"No content in '{Path(p).name}'")
                    continue
                chunks = self.splitter.split_documents(raw)
                all_chunks.extend(chunks)

                file_name = Path(p).name
                file_ext = Path(p).suffix.lower().lstrip(".")
                loaded_meta.append({
                    "filename": file_name,
                    "ext": file_ext,
                    "chunks": len(chunks),
                    "size_kb": round(Path(p).stat().st_size / 1024, 1)
                })
            except Exception as e:
                errors.append(f"'{Path(p).name}': {str(e)}")

        return all_chunks, loaded_meta, errors
